import requests
import sys
from bs4 import BeautifulSoup
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Alignment, Font

#https://www.whatismybrowser.com/detect/what-http-headers-is-my-browser-sending/
headers = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/131.0.0.0 Safari/537.36"
}

season = int(sys.argv[1])

# dictionary with team fixtures
premierLeague = {
                "Arsenal" : f"https://www.transfermarkt.com/fc-arsenal/spielplan/verein/11/saison_id/{season}",
                "Chelsea" : f"https://www.transfermarkt.com/fc-chelsea/spielplan/verein/631/saison_id/{season}",
                "Liverpool" : f"https://www.transfermarkt.com/fc-liverpool/spielplan/verein/31/saison_id/{season}", 
                "ManchesterCity" : f"https://www.transfermarkt.com/manchester-city/spielplan/verein/281/saison_id/{season}",
                "ManchesterUnited" : f"https://www.transfermarkt.com/manchester-united/spielplan/verein/985/saison_id/{season}"
                }

laLiga = {
        "AtleticoMadrid" : f"https://www.transfermarkt.com/atletico-madrid/spielplan/verein/13/saison_id/{season}", 
        "Barcelona" : f"https://www.transfermarkt.com/fc-barcelona/spielplan/verein/131/saison_id/{season}",
        "RealMadrid" : f"https://www.transfermarkt.com/real-madrid/spielplan/verein/418/saison_id/{season}"
        }

bundesLiga = {
            "BayerLeverkusen" : f"https://www.transfermarkt.com/bayer-04-leverkusen/spielplan/verein/15/saison_id/{season}",
            "BayernMunchen" : f"https://www.transfermarkt.com/fc-bayern-munchen/spielplan/verein/27/saison_id/{season}", 
            "BorussiaDortmund" : f"https://www.transfermarkt.com/borussia-dortmund/spielplan/verein/16/saison_id/{season}", 
            "Frankfurt" : f"https://www.transfermarkt.com/eintracht-frankfurt/spielplan/verein/24/saison_id/{season}",
            "Mainz" : f"https://www.transfermarkt.com/1-fsv-mainz-05/spielplan/verein/39/saison_id/{season}",
            "RBLeipzig" : f"https://www.transfermarkt.com/rasenballsport-leipzig/spielplan/verein/23826/saison_id/{season}",
            "Stuttgart" : f"https://www.transfermarkt.com/vfb-stuttgart/spielplan/verein/79/saison_id/{season}"
            }

serieA = {
    "ACMilan" : f"https://www.transfermarkt.com/ac-mailand/spielplan/verein/5/saison_id/{season}", 
    "InterMilan" : f"https://www.transfermarkt.com/inter-mailand/spielplan/verein/46/saison_id/{season}",
    "Juventus" : f"https://www.transfermarkt.com/juventus-turin/spielplan/verein/506/saison_id/{season}",
    "Lazio" : f"https://www.transfermarkt.com/lazio-rom/spielplan/verein/398/saison_id/{season}",
    "Napoli" : f"https://www.transfermarkt.com/ssc-neapel/spielplan/verein/6195/saison_id/{season}", 
    "Roma" : f"https://www.transfermarkt.com/as-rom/spielplan/verein/12/saison_id/{season}"
        }   

ligue1 = {
        "Lille" : f"https://www.transfermarkt.com/losc-lille/spielplan/verein/1082/saison_id/{season}",
        "Lyon" : f"https://www.transfermarkt.com/olympique-lyon/spielplan/verein/1041/saison_id/{season}",
        "Marseille" : f"https://www.transfermarkt.com/olympique-marseille/spielplan/verein/244/saison_id/{season}",
        "Monaco" : f"https://www.transfermarkt.com/as-monaco/spielplan/verein/162/saison_id/{season}",
        "Nice" : f"https://www.transfermarkt.com/ogc-nizza/spielplan/verein/417/saison_id/{season}",
        "PSG" : f"https://www.transfermarkt.com/fc-paris-saint-germain/spielplan/verein/583/saison_id/{season}"
        }

leagues ={
        "England" : premierLeague,
        "Spain" : laLiga,
        "Germany" : bundesLiga,
        "Italy" : serieA,
        "France" : ligue1
        }

# generate new Excel sheet
file_path_detailed = f"team_stats_{season}_detailed.xlsx"
file_path = f"team_stats_{season}.xlsx"
sheet_name = "Generated"
wb = Workbook()
wb.create_sheet(sheet_name)
wb.remove(wb["Sheet"])
sheet=wb[sheet_name]
sheet.freeze_panes = "C2"
wb.save(file_path_detailed)
wb.save(file_path)

# function to retrieve data from TransferMarkt
def retrieve_data(index, link, sect):
    response = requests.get(link, headers=headers)
    response.status_code
    soup = BeautifulSoup(response.content, "html.parser")

    section = soup.find('a', href=sect)
    table = section.find_parent('thead').find_next_sibling('tbody')
    data = [td.get_text(strip=True) for td in table.find_all('td', class_='zentriert')]

    for i, value in enumerate(data[:12]):
        if(value == '-'):
            data[i] = 0

    # compute each field
    homeMatches = int(data[0])
    homeWins = int(data[1])
    homeDraws = int(data[2])
    homeLosses = int(data[3])
    homePoints = round(float(data[4]), 2)

    if homeMatches == 0:
        homeWins_p = 0
        homeDraws_p = 0
        homeLosses_p = 0
        homeGoals_f = 0
        homeGoals_a = 0
        homeGoals_f_v = 0
        homeGoals_a_v = 0
    else:
        homeWins_p = round(homeWins / homeMatches * 100, 2)
        homeDraws_p = round(homeDraws / homeMatches * 100, 2)
        homeLosses_p = round(homeLosses / homeMatches* 100 , 2)
        homeGoals = data[5].split(":")
        homeGoals_f = int(homeGoals[0])
        homeGoals_a = int(homeGoals[1])
        homeGoals_f_v = round(homeGoals_f / homeMatches , 2)
        homeGoals_a_v = round(homeGoals_a / homeMatches, 2)
    homeGD = homeGoals_f - homeGoals_a

    awayMatches = int(data[6])
    awayWins = int(data[7])
    awayDraws = int(data[8])
    awayLosses = int(data[9])
    awayPoints = round(float(data[10]), 2)

    if awayMatches == 0:
        awayWins_p = 0
        awayDraws_p = 0
        awayLosses_p = 0
        awayGoals_f = 0
        awayGoals_a = 0
        awayGoals_f_v = 0
        awayGoals_a_v = 0
    else:
        awayWins_p = round(awayWins / awayMatches * 100, 2)
        awayDraws_p = round(awayDraws / awayMatches * 100, 2)
        awayLosses_p = round(awayLosses / awayMatches* 100 , 2)
        awayGoals = data[11].split(":")
        awayGoals_f = int(awayGoals[0])
        awayGoals_a = int(awayGoals[1])
        awayGoals_f_v = round(awayGoals_f / awayMatches, 2)
        awayGoals_a_v = round(awayGoals_a / awayMatches, 2)
    awayGD = awayGoals_f - awayGoals_a

    totalMatches = homeMatches + awayMatches
    totalWins = homeWins + awayWins
    totalDraws = homeDraws + awayDraws
    totalLosses = homeLosses + awayLosses
    totalGoals_f = homeGoals_f + awayGoals_f
    totalGoals_a = homeGoals_a + awayGoals_a
    if totalMatches == 0:
        totalWins_p = 0
        totalDraws_p = 0
        totalLosses_p = 0
        totalPoints = 0
        totalGoals_f_v = 0
        totalGoals_a_v = 0
    else:
        totalWins_p = round(totalWins / totalMatches * 100, 2)
        totalDraws_p = round(totalDraws / totalMatches * 100 , 2)
        totalLosses_p = round(totalLosses / totalMatches * 100, 2)
        totalPoints = round((3 * totalWins + totalDraws) / totalMatches, 2)
        totalGoals_f_v = round(totalGoals_f / totalMatches, 2)
        totalGoals_a_v = round(totalGoals_a / totalMatches, 2)
    totalGD = totalGoals_f - totalGoals_a

    #update detailed Excel
    wb = load_workbook(file_path_detailed)
    sheet = wb[sheet_name]

    sheet[f"C{index}"] = homeMatches
    sheet[f"D{index}"] = homeWins
    sheet[f"E{index}"] = homeWins_p
    sheet[f"F{index}"] = homeDraws
    sheet[f"G{index}"] = homeDraws_p
    sheet[f"H{index}"] = homeLosses
    sheet[f"I{index}"] = homeLosses_p
    sheet[f"J{index}"] = homePoints
    sheet[f"K{index}"] = homeGoals_f
    sheet[f"L{index}"] = homeGoals_f_v
    sheet[f"M{index}"] = homeGoals_a
    sheet[f"N{index}"] = homeGoals_a_v
    sheet[f"O{index}"] = homeGD

    sheet[f"P{index}"] = awayMatches
    sheet[f"Q{index}"] = awayWins
    sheet[f"R{index}"] = awayWins_p
    sheet[f"S{index}"] = awayDraws
    sheet[f"T{index}"] = awayDraws_p
    sheet[f"U{index}"] = awayLosses
    sheet[f"V{index}"] = awayLosses_p
    sheet[f"W{index}"] = awayPoints
    sheet[f"X{index}"] = awayGoals_f
    sheet[f"Y{index}"] = awayGoals_f_v
    sheet[f"Z{index}"] = awayGoals_a
    sheet[f"AA{index}"] = awayGoals_a_v
    sheet[f"AB{index}"] = awayGD

    sheet[f"AC{index}"] = totalMatches
    sheet[f"AD{index}"] = totalWins
    sheet[f"AE{index}"] = totalWins_p
    sheet[f"AF{index}"] = totalDraws
    sheet[f"AG{index}"] = totalDraws_p
    sheet[f"AH{index}"] = totalLosses
    sheet[f"AI{index}"] = totalLosses_p
    sheet[f"AJ{index}"] = totalPoints
    sheet[f"AK{index}"] = totalGoals_f
    sheet[f"AL{index}"] = totalGoals_f_v
    sheet[f"AM{index}"] = totalGoals_a
    sheet[f"AN{index}"] = totalGoals_a_v
    sheet[f"AO{index}"] = totalGD

    wb.save(file_path_detailed)

    #update Excel
    wb = load_workbook(file_path)
    sheet = wb[sheet_name]

    sheet[f"C{index}"] = homeMatches
    sheet[f"D{index}"] = homeWins
    sheet[f"E{index}"] = homeDraws
    sheet[f"F{index}"] = homeLosses
    sheet[f"G{index}"] = homeGD

    sheet[f"H{index}"] = awayMatches
    sheet[f"I{index}"] = awayWins
    sheet[f"J{index}"] = awayDraws
    sheet[f"K{index}"] = awayLosses
    sheet[f"L{index}"] = awayGD

    sheet[f"M{index}"] = totalMatches
    sheet[f"N{index}"] = totalWins
    sheet[f"O{index}"] = totalDraws
    sheet[f"P{index}"] = totalLosses
    sheet[f"Q{index}"] = totalGD

    wb.save(file_path)


# call function based on league
index = 2
sect = ""

for k,v in leagues.items():
    for i in v:
        if k == "England":
            sect = f"/premier-league/startseite/wettbewerb/GB1/saison_id/{season}"
        elif k == "Spain":
            sect = f"/laliga/startseite/wettbewerb/ES1/saison_id/{season}"
            if index == 2 + len(premierLeague):
                index +=2
        elif k == "Germany":
            sect = f"/bundesliga/startseite/wettbewerb/L1/saison_id/{season}"
            if index == 2 + len(premierLeague) + 2 + len(laLiga):
                index += 2
        elif k == "Italy":
            sect = f"/serie-a/startseite/wettbewerb/IT1/saison_id/{season}"
            if index == 2 + len(premierLeague) + 2 + len(laLiga) + 2 + len(bundesLiga):
                index += 2
        elif k == "France":
            sect = f"/ligue-1/startseite/wettbewerb/FR1/saison_id/{season}"
            if index == 2 + len(premierLeague) + 2 + len(laLiga) + 2 + len(bundesLiga) + 2 + len(serieA):
                index += 2
        retrieve_data(index, v[i], sect)
        index +=1

# populate detailed Excel sheet with team names and metrics
wb = load_workbook(file_path_detailed)
sheet = wb[sheet_name]

for row in range(1,sheet.max_row + 1):
    for col in range(1, sheet.max_column + 1):
        sheet.cell(row,col).font = Font(name='Helvetica', size=12, bold=True, color = '000000')
        sheet.cell(row,col).alignment = Alignment(horizontal='center', vertical='center')
        
sheet["A1"] = "Premier League"
sheet["B1"] = "Season"
sheet["C1"] = "Home matches"
sheet["D1"] = "Home wins"
sheet["E1"] = "%"
sheet["F1"] = "Home draws"
sheet["G1"] = "%"
sheet["H1"] = "Home losses"
sheet["I1"] = "%"
sheet["J1"] = "Home points / match"
sheet["K1"] = "Home goals scored"
sheet["L1"] = " / match"
sheet["M1"] = "Home goals conceded"
sheet["N1"] = " / match"
sheet["O1"] = "Home goal difference"
sheet["P1"] = "Away matches"
sheet["Q1"] = "Away wins"
sheet["R1"] = "%"
sheet["S1"] = "Away draws"
sheet["T1"] = "%"
sheet["U1"] = "Aways losses"
sheet["V1"] = "%"
sheet["W1"] = "Aways points / match"
sheet["X1"] = "Aways goals scored"
sheet["Y1"] = " / match"
sheet["Z1"] = "Away goals conceded"
sheet["AA1"] = " / match"
sheet["AB1"] = "Away home difference"
sheet["AC1"] = "Total matches"
sheet["AD1"] = "Total wins"
sheet["AE1"] = "%"
sheet["AF1"] = "Total draws"
sheet["AG1"] = "%"
sheet["AH1"] = "Total losses"
sheet["AI1"] = "%"
sheet["AJ1"] = "Total points / match"
sheet["AK1"] = "Total goals scored"
sheet["AL1"] = " / match"
sheet["AM1"] = "Total goals conceded"
sheet["AN1"] = " / match"
sheet["AO1"] = "Total goal difference"

index = 2

for k,v in leagues.items():
    for i in v:
        if k == "Spain":
            if index == 2 + len(premierLeague):
                sheet[f"A{index+1}"] = "La Liga"
                index +=2
        elif k == "Germany":
            if index == 2 + len(premierLeague) + 2 + len(laLiga):
                sheet[f"A{index+1}"] = "Bundesliga"
                index += 2
        elif k == "Italy":
            if index == 2 + len(premierLeague) + 2 + len(laLiga) + 2 + len(bundesLiga):
                sheet[f"A{index+1}"] = "Serie A"
                index += 2
        elif k == "France":
            if index == 2 + len(premierLeague) + 2 + len(laLiga) + 2 + len(bundesLiga) + 2 + len(serieA):
                sheet[f"A{index+1}"] = "Ligue 1"
                index += 2
        sheet[f"A{index}"] = i
        sheet[f"B{index}"] = str(season) + "/" + str(season+1)
        index +=1

for col in sheet.columns:
     max_width = 0
     column = col[0].column_letter
     for cell in col:
             if len(str(cell.value)) > max_width:
                 max_width = len(str(cell.value))
     set_col_width = max_width + 7
     sheet.column_dimensions[column].width = set_col_width
     
wb.save(file_path_detailed)

# populate Excel sheet with team names and metrics
wb = load_workbook(file_path)
sheet = wb[sheet_name]

for row in range(1,sheet.max_row + 1):
    for col in range(1, sheet.max_column + 1):
        sheet.cell(row,col).font = Font(name='Helvetica', size=12, bold=True, color = '000000')
        sheet.cell(row,col).alignment = Alignment(horizontal='center', vertical='center')

sheet["A1"] = "Premier League"
sheet["B1"] = "Season"
sheet["C1"] = "Home matches"
sheet["D1"] = "Home wins"
sheet["E1"] = "Home draws"
sheet["F1"] = "Home losses"
sheet["G1"] = "Home goal difference"
sheet["H1"] = "Away matches"
sheet["I1"] = "Away wins"
sheet["J1"] = "Away draws"
sheet["K1"] = "Aways losses"
sheet["L1"] = "Away goal difference"
sheet["M1"] = "Total matches"
sheet["N1"] = "Total wins"
sheet["O1"] = "Total draws"
sheet["P1"] = "Total losses"
sheet["Q1"] = "Total goal difference"

index = 2

for k,v in leagues.items():
    for i in v:
        if k == "Spain":
            if index == 2 + len(premierLeague):
                sheet[f"A{index+1}"] = "La Liga"
                index +=2
        elif k == "Germany":
            if index == 2 + len(premierLeague) + 2 + len(laLiga):
                sheet[f"A{index+1}"] = "Bundesliga"
                index += 2
        elif k == "Italy":
            if index == 2 + len(premierLeague) + 2 + len(laLiga) + 2 + len(bundesLiga):
                sheet[f"A{index+1}"] = "Serie A"
                index += 2
        elif k == "France":
            if index == 2 + len(premierLeague) + 2 + len(laLiga) + 2 + len(bundesLiga) + 2 + len(serieA):
                sheet[f"A{index+1}"] = "Ligue 1"
                index += 2
        sheet[f"A{index}"] = i
        sheet[f"B{index}"] = str(season) + "/" + str(season+1)
        index +=1

for col in sheet.columns:
     max_width = 0
     column = col[0].column_letter
     for cell in col:
             if len(str(cell.value)) > max_width:
                 max_width = len(str(cell.value))
     set_col_width = max_width + 7
     sheet.column_dimensions[column].width = set_col_width
     
wb.save(file_path)
